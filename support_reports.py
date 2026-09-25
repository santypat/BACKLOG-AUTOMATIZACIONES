"""Generación del reporte descargable del historial de soportes."""

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from backlog_domain import interpretar_nombre_soporte


REPORT_COLUMNS = (
    "ID",
    "Nombre del soporte",
    "Origen",
    "Desarrollador",
    "Célula",
    "Estado",
    "Tipo de soporte",
    "Prioridad",
    "Horas empleadas",
    "Fecha de ingreso",
    "Fecha de entrega",
    "Descripción",
    "Observaciones",
)


def _texto_excel_seguro(valor):
    """Evita que contenido ingresado por usuarios se ejecute como fórmula."""
    if valor is None or pd.isna(valor):
        return ""
    texto = str(valor)
    if texto.startswith(("=", "+", "-", "@")):
        return f"'{texto}"
    return texto


def preparar_historial_soportes(soportes_df):
    """Construye la tabla legible que se incluirá en el archivo Excel."""
    registros = []

    for _, soporte in soportes_df.iterrows():
        nombre, es_automatizacion = interpretar_nombre_soporte(
            soporte.get("desarrollo")
        )
        registros.append({
            "ID": soporte.get("id"),
            "Nombre del soporte": _texto_excel_seguro(nombre),
            "Origen": (
                "Automatización"
                if es_automatizacion
                else "Soporte independiente"
            ),
            "Desarrollador": _texto_excel_seguro(
                soporte.get("desarrollador")
            ),
            "Célula": _texto_excel_seguro(soporte.get("celula")),
            "Estado": _texto_excel_seguro(soporte.get("estado")),
            "Tipo de soporte": _texto_excel_seguro(
                soporte.get("tipo_soporte")
            ),
            "Prioridad": _texto_excel_seguro(soporte.get("prioridad")),
            "Horas empleadas": soporte.get("horas_empleadas"),
            "Fecha de ingreso": pd.to_datetime(
                soporte.get("fecha_ingreso"),
                errors="coerce",
            ),
            "Fecha de entrega": pd.to_datetime(
                soporte.get("fecha_entrega"),
                errors="coerce",
            ),
            "Descripción": _texto_excel_seguro(soporte.get("descripcion")),
            "Observaciones": _texto_excel_seguro(
                soporte.get("observaciones")
            ),
        })

    return pd.DataFrame(registros, columns=REPORT_COLUMNS)


def generar_reporte_soportes_excel(soportes_df):
    """Devuelve un XLSX con el historial completo de soportes."""
    reporte = preparar_historial_soportes(soportes_df)
    buffer = io.BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        reporte.to_excel(
            writer,
            index=False,
            sheet_name="Historial de soportes",
        )

        hoja = writer.sheets["Historial de soportes"]
        hoja.freeze_panes = "A2"
        hoja.sheet_view.showGridLines = False

        relleno_encabezado = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        for celda in hoja[1]:
            celda.fill = relleno_encabezado
            celda.font = Font(color="FFFFFF", bold=True)
            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        anchos = {
            "A": 10,
            "B": 34,
            "C": 23,
            "D": 30,
            "E": 40,
            "F": 16,
            "G": 18,
            "H": 14,
            "I": 17,
            "J": 18,
            "K": 18,
            "L": 55,
            "M": 55,
        }
        for columna, ancho in anchos.items():
            hoja.column_dimensions[columna].width = ancho

        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                celda.alignment = Alignment(
                    vertical="top",
                    wrap_text=celda.column in (2, 4, 5, 12, 13),
                )
            fila[8].number_format = "0.00"
            fila[9].number_format = "dd/mm/yyyy"
            fila[10].number_format = "dd/mm/yyyy"

        if hoja.max_row >= 2:
            tabla = Table(
                displayName="HistorialSoportes",
                ref=f"A1:M{hoja.max_row}",
            )
            tabla.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            hoja.add_table(tabla)

    return buffer.getvalue()
