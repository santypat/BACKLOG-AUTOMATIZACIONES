import io
import unittest

import pandas as pd
from openpyxl import load_workbook

from support_reports import (
    REPORT_COLUMNS,
    generar_reporte_soportes_excel,
    preparar_historial_soportes,
)


class ReporteSoportesTest(unittest.TestCase):
    def setUp(self):
        self.soportes = pd.DataFrame([
            {
                "id": 10,
                "desarrollo": "Robot de novedades",
                "desarrollador": "Ana Ruiz",
                "celula": "TECNOLOGIA",
                "estado": "Finalizado",
                "tipo_soporte": "Mejora",
                "prioridad": "MEDIA",
                "horas_empleadas": 4.5,
                "fecha_ingreso": "2026-09-17",
                "fecha_entrega": "2026-09-24",
                "descripcion": "Primera línea\nSegunda línea",
                "observaciones": "=2+2",
            },
            {
                "id": 11,
                "desarrollo": "__SOPORTE_INDEPENDIENTE__:Equipo nuevo",
                "desarrollador": "Luis Pérez",
                "celula": "TRANSVERSALES",
                "estado": "Descartado",
                "tipo_soporte": "Ajuste",
                "prioridad": "BAJA",
                "horas_empleadas": 0,
                "fecha_ingreso": "2026-09-25",
                "fecha_entrega": "2026-09-25",
                "descripcion": "No requerido",
                "observaciones": "",
            },
        ])

    def test_prepara_todos_los_estados_y_origenes(self):
        reporte = preparar_historial_soportes(self.soportes)
        self.assertEqual(tuple(reporte.columns), REPORT_COLUMNS)
        self.assertEqual(len(reporte), 2)
        self.assertEqual(reporte.iloc[1]["Origen"], "Soporte independiente")
        self.assertEqual(reporte.iloc[1]["Estado"], "Descartado")
        self.assertEqual(reporte.iloc[0]["Observaciones"], "'=2+2")

    def test_genera_excel_legible(self):
        contenido = generar_reporte_soportes_excel(self.soportes)
        libro = load_workbook(io.BytesIO(contenido), data_only=False)
        hoja = libro["Historial de soportes"]

        self.assertEqual(hoja.freeze_panes, "A2")
        self.assertEqual(hoja.max_row, 3)
        self.assertEqual(hoja.max_column, len(REPORT_COLUMNS))
        self.assertEqual(hoja["F3"].value, "Descartado")
        self.assertEqual(hoja["J2"].number_format, "dd/mm/yyyy")
        self.assertEqual(hoja["M2"].data_type, "s")


if __name__ == "__main__":
    unittest.main()
